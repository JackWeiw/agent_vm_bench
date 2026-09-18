"""Unit tests for the oversub cross-architecture compare tool.

Pure helpers in ``bench_core.compare_sweeps`` are module-level so pytest
imports them directly. The CSVs are synthetic (no live sweep runs): they match
the driver's ``trajectory-detail.csv`` / ``trial-summary.csv`` contract
(column names from ``bench_core.oversub.TRAJECTORY_COLUMNS`` /
``TRIAL_COLUMNS``), so a real sweep's output is consumable unchanged.
"""
from __future__ import annotations

import csv
from pathlib import Path

import pytest
import yaml

from bench_core.compare_sweeps import (
    build_ratio_summary,
    build_tidy,
    build_trajectory_delta,
    check_n_parity,
    load_manifest,
    main,
    run_comparison,
)
from bench_core.oversub import TRAJECTORY_COLUMNS, TRIAL_COLUMNS, load_sweep_config

REPO = Path(__file__).resolve().parents[3]


# ---- synthetic CSV fixtures -----------------------------------------------


def _write_traj(path: Path, rows: list[dict]) -> None:
    """Write a trajectory-detail.csv; missing fields -> blank (restval)."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=TRAJECTORY_COLUMNS, extrasaction="ignore", restval="")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def _write_trial(path: Path, rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=TRIAL_COLUMNS, extrasaction="ignore", restval="")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def _traj_row(mode: str, ratio: int, traj_id: str, elapsed: float) -> dict:
    """One trajectory-detail row; breakdown cols derived from elapsed so the
    component heatmaps have non-degenerate data."""
    return {
        "mode": mode,
        "ratio": ratio,
        "repeat": 1,
        "trajectory_id": traj_id,
        "sandbox_index": 0,
        "n_steps": 3,
        "n_failed": 0,
        "n_timeout": 0,
        "success_rate": 1.0,
        "elapsed_sec": elapsed,
        "slice_total_sec": elapsed * 0.9,
        "exec_sec": elapsed * 0.5,
        "resume_sec": elapsed * 0.1,
        "pause_sec": elapsed * 0.1,
        "requested_delay_sec": elapsed * 0.05,
        "create_sec": 0.05,
        "kill_sec": 0.05,
        "interaction_total_sec": elapsed * 0.8,
        "slot_contention_wait_sec": elapsed * 0.1,
        "natural_delay_sec": elapsed * 0.05,
        "capacity_wait_sec": 0.0,
        "rate_pacing_wait_sec": 0.0,
        "inflight_wait_sec": 0.0,
        "resume_rate_pacing_wait_sec": 0.0,
        "pause_rate_pacing_wait_sec": 0.0,
        "resume_inflight_wait_sec": 0.0,
        "pause_inflight_wait_sec": 0.0,
        "running_slot_held_sec": elapsed * 0.9,
        "create_error_type": "",
        "kill_error_type": "",
    }


def _trial_row(mode: str, ratio: int, *, n: int, wall: float, valid: bool = True) -> dict:
    return {
        "mode": mode,
        "ratio": ratio,
        "oversubscription_ratio": ratio,
        "repeat": 1,
        "running_concurrency": n,
        "target_count": ratio * n,
        "total": ratio * n,
        "succeeded": ratio * n,
        "failed": 0,
        "failure_rate": 0.0,
        "peak_active": n,
        "granted": n,
        "avg_queue_wait_sec": 0.0,
        "control_dispatched": n,
        "test_duration": 600,
        "wall_sec": wall,
        "tasks_per_sec": (ratio * n) / wall if wall else 0.0,
        "steps_per_sec": 0.01,
        "lifecycle_overhead_pct": 10.0,
        "return_code": 0 if valid else 1,
        "valid": valid,
        "reused": False,
        "started_at": "",
        "completed_at": "",
        "error": "",
        "trial_dir": "",
        "run_summary_path": "",
    }


def _seed_series(
    root: Path, *, ratios: list[int], elapsed: dict[tuple[int, str], float], n: int = 384, valid: bool = True
) -> Path:
    """Write a series dir with trajectory-detail.csv + trial-summary.csv.

    elapsed: {(ratio, trajectory_id): seconds}. trial wall_sec = the per-ratio
    max elapsed (any positive number; the trial-derived metrics are not the
    focus of the assertion math, just non-degenerate)."""
    root.mkdir(parents=True, exist_ok=True)
    traj_rows = [_traj_row("lifecycle", r, t, e) for (r, t), e in elapsed.items()]
    _write_traj(root / "trajectory-detail.csv", traj_rows)
    wall = {r: max((e for (rr, _), e in elapsed.items() if rr == r), default=10.0) for r in ratios}
    trial_rows = [_trial_row("lifecycle", r, n=n, wall=wall[r], valid=valid) for r in ratios]
    _write_trial(root / "trial-summary.csv", trial_rows)
    return root


# ---- manifest load --------------------------------------------------------


def test_load_manifest_valid(tmp_path):
    arm = _seed_series(
        tmp_path / "arm", ratios=[1, 2], elapsed={(1, "t0"): 2.0, (1, "t1"): 4.0, (2, "t0"): 4.0, (2, "t1"): 8.0}
    )
    x86 = _seed_series(
        tmp_path / "x86", ratios=[1, 2], elapsed={(1, "t0"): 1.0, (1, "t1"): 2.0, (2, "t0"): 2.0, (2, "t1"): 4.0}
    )
    m_path = tmp_path / "m.yaml"
    m_path.write_text(
        yaml.safe_dump(
            {
                "baseline_series": "arm",
                "series": [
                    {"label": "arm", "arch": "arm", "dir": str(arm), "caps": {"l3_mb": 256}},
                    {"label": "x86", "arch": "x86", "dir": str(x86), "caps": {"freq_ghz": 3.0, "l3_mb": 64}},
                ],
            }
        ),
        encoding="utf-8",
    )
    m = load_manifest(m_path)
    assert m["baseline_series"] == "arm"
    assert [s["label"] for s in m["series"]] == ["arm", "x86"]
    assert m["series"][0]["caps"] == {"l3_mb": 256}


def test_load_manifest_rejects_unknown_key(tmp_path):
    m_path = tmp_path / "m.yaml"
    m_path.write_text(
        yaml.safe_dump({"baseline_series": "a", "serie": []}),  # typo
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="unknown manifest key"):
        load_manifest(m_path)


def test_load_manifest_rejects_unknown_baseline(tmp_path):
    m_path = tmp_path / "m.yaml"
    m_path.write_text(
        yaml.safe_dump({"baseline_series": "ghost", "series": [{"label": "arm", "dir": str(tmp_path)}]}),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="baseline_series"):
        load_manifest(m_path)


def test_load_manifest_rejects_duplicate_label(tmp_path):
    m_path = tmp_path / "m.yaml"
    m_path.write_text(
        yaml.safe_dump(
            {
                "baseline_series": "arm",
                "series": [
                    {"label": "arm", "dir": str(tmp_path)},
                    {"label": "arm", "dir": str(tmp_path)},
                ],
            }
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="duplicate series label"):
        load_manifest(m_path)


def test_shipped_lifecycle_1to6_sweep_config_loads():
    """lifecycle-1to6.yaml must load via the sweep-config known-keys check
    (every key must be a real oversub-bench knob, not a manifest key)."""
    cfg = load_sweep_config(REPO / "config" / "oversub" / "lifecycle-1to6.yaml")
    assert cfg["base_config"] == "config/common/replay.yaml"
    assert cfg["ratios"] == [1, 2, 3, 4, 5, 6]
    assert cfg["modes"] == ["lifecycle"]
    assert cfg["running_concurrency"] == 384  # pinned so both archs share N
    assert cfg["repeats"] == 3


def test_shipped_compare_manifest_loads():
    """The shipped manifest template loads structurally (paths resolve even
    though the dirs do not exist on disk -- load_manifest does not stat them)."""
    m = load_manifest(REPO / "config" / "oversub" / "compare-manifest.yaml")
    assert m["baseline_series"] == "arm-ampere"
    labels = [s["label"] for s in m["series"]]
    assert "arm-ampere" in labels and "x86-base" in labels


# ---- tidy + ratio summary -------------------------------------------------


def _two_series(tmp_path, *, x86_n: int = 384):
    """arm baseline (slower) + x86 (faster, degrades worse). Clean numbers."""
    arm = _seed_series(
        tmp_path / "arm",
        ratios=[1, 2],
        elapsed={(1, "t0"): 2.0, (1, "t1"): 4.0, (2, "t0"): 4.0, (2, "t1"): 8.0},  # medians 3.0, 6.0
    )
    x86 = _seed_series(
        tmp_path / "x86",
        ratios=[1, 2],
        n=x86_n,
        elapsed={(1, "t0"): 1.0, (1, "t1"): 2.0, (2, "t0"): 3.0, (2, "t1"): 5.0},  # medians 1.5, 4.0
    )
    return {
        "baseline_series": "arm",
        "series": [
            {"label": "arm", "arch": "arm", "dir": str(arm), "caps": {"l3_mb": 256}},
            {"label": "x86", "arch": "x86", "dir": str(x86), "caps": {"freq_ghz": 3.0, "l3_mb": 64}},
        ],
    }


def test_build_tidy_shape_and_caps(tmp_path):
    manifest = _two_series(tmp_path)
    from bench_core.compare_sweeps import _read_series_tables

    tables = _read_series_tables(manifest)
    tidy = build_tidy(manifest, tables)
    # 2 series * 2 ratios * 2 trajectories * 19 metrics = 152 rows.
    assert len(tidy) == 2 * 2 * 2 * 19
    # caps exploded to columns (union: freq_ghz, l3_mb); arm lacks freq_ghz -> NaN.
    assert {"series", "arch", "l3_mb", "freq_ghz", "mode", "ratio", "trajectory_id", "metric", "value"} <= set(
        tidy.columns
    )
    arm_rows = tidy[tidy["series"] == "arm"]
    assert arm_rows["l3_mb"].iloc[0] == 256
    assert arm_rows["freq_ghz"].isna().all()


def test_build_ratio_summary_medians_and_deltas(tmp_path):
    manifest = _two_series(tmp_path)
    from bench_core.compare_sweeps import _read_series_tables

    tables = _read_series_tables(manifest)
    tidy = build_tidy(manifest, tables)
    rs = build_ratio_summary(tidy, tables, manifest)

    def get(series, ratio, metric):
        row = rs[(rs["series"] == series) & (rs["ratio"] == ratio) & (rs["metric"] == metric)]
        assert len(row) == 1, f"expected one row for {series}/{ratio}/{metric}, got {len(row)}"
        return row.iloc[0]

    # Trajectory-derived medians.
    assert get("arm", 1, "elapsed_sec")["median_sec"] == 3.0
    assert get("arm", 2, "elapsed_sec")["median_sec"] == 6.0
    assert get("x86", 1, "elapsed_sec")["median_sec"] == 1.5
    assert get("x86", 2, "elapsed_sec")["median_sec"] == 4.0

    # delta vs baseline (arm).
    r = get("x86", 1, "elapsed_sec")
    assert r["delta_vs_baseline_sec"] == pytest.approx(1.5 - 3.0)
    assert r["pct_delta_vs_baseline"] == pytest.approx((1.5 - 3.0) / 3.0 * 100)
    # arm's own delta vs baseline is 0.
    assert get("arm", 1, "elapsed_sec")["delta_vs_baseline_sec"] == 0.0

    # within-series degradation vs k=1.
    assert get("arm", 1, "elapsed_sec")["degradation_vs_ratio1_pct"] == 0.0  # k=1 by definition
    assert get("arm", 2, "elapsed_sec")["degradation_vs_ratio1_pct"] == pytest.approx((6.0 - 3.0) / 3.0 * 100)
    assert get("x86", 2, "elapsed_sec")["degradation_vs_ratio1_pct"] == pytest.approx(round((4.0 - 1.5) / 1.5 * 100, 3))
    # x86 degrades worse (~167%) than arm (100%) -- the cross-arch insight.
    assert (
        get("x86", 2, "elapsed_sec")["degradation_vs_ratio1_pct"]
        > get("arm", 2, "elapsed_sec")["degradation_vs_ratio1_pct"]
    )


def test_build_ratio_summary_no_k1_blanks_degradation(tmp_path):
    """A series with no k=1 trial -> degradation blank, not a crash."""
    arm = _seed_series(tmp_path / "arm", ratios=[1, 2], elapsed={(1, "t0"): 2.0, (2, "t0"): 4.0})
    x86 = _seed_series(tmp_path / "x86", ratios=[2], elapsed={(2, "t0"): 3.0})  # no k=1
    manifest = {
        "baseline_series": "arm",
        "series": [
            {"label": "arm", "arch": "arm", "dir": str(arm), "caps": {}},
            {"label": "x86", "arch": "x86", "dir": str(x86), "caps": {}},
        ],
    }
    from bench_core.compare_sweeps import _read_series_tables

    tables = _read_series_tables(manifest)
    tidy = build_tidy(manifest, tables)
    rs = build_ratio_summary(tidy, tables, manifest)
    row = rs[(rs["series"] == "x86") & (rs["ratio"] == 2) & (rs["metric"] == "elapsed_sec")].iloc[0]
    # No k=1 baseline for x86 -> degradation is NaN (not a crash, not 0).
    import math

    assert math.isnan(row["degradation_vs_ratio1_pct"])


# ---- per-trajectory delta -------------------------------------------------


def test_build_trajectory_delta_aligned_and_orphan(tmp_path):
    """t0 aligns across series; t1 is an orphan (present in arm, absent in x86)
    -> x86 columns blank + delta NaN."""
    arm = _seed_series(
        tmp_path / "arm",
        ratios=[1],
        elapsed={(1, "t0"): 2.0, (1, "t1"): 5.0},
    )
    x86 = _seed_series(
        tmp_path / "x86",
        ratios=[1],
        elapsed={(1, "t0"): 1.5},  # no t1
    )
    manifest = {
        "baseline_series": "arm",
        "series": [
            {"label": "arm", "arch": "arm", "dir": str(arm), "caps": {}},
            {"label": "x86", "arch": "x86", "dir": str(x86), "caps": {}},
        ],
    }
    from bench_core.compare_sweeps import _read_series_tables

    tables = _read_series_tables(manifest)
    td = build_trajectory_delta(tables, manifest)
    by_id = {r["trajectory_id"]: r for _, r in td.iterrows()}
    # t0 aligns: arm=2.0, x86=1.5, delta=-0.5.
    assert by_id["t0"]["elapsed_sec_baseline"] == 2.0
    assert by_id["t0"]["elapsed_sec__x86"] == 1.5
    assert by_id["t0"]["delta_sec__x86"] == pytest.approx(-0.5)
    # t1 orphan: arm has it (5.0), x86 blank -> delta NaN.
    assert by_id["t1"]["elapsed_sec_baseline"] == 5.0
    assert pd_isna(by_id["t1"]["elapsed_sec__x86"])
    assert pd_isna(by_id["t1"]["delta_sec__x86"])


def pd_isna(v) -> bool:
    import pandas as pd

    return bool(pd.isna(v))


# ---- N-parity -------------------------------------------------------------


def test_check_n_parity_matched(tmp_path):
    manifest = _two_series(tmp_path, x86_n=384)
    from bench_core.compare_sweeps import _read_series_tables

    ok, warns = check_n_parity(_read_series_tables(manifest))
    assert ok is True
    assert all("MISMATCH" not in w for w in warns)


def test_check_n_parity_mismatched(tmp_path):
    manifest = _two_series(tmp_path, x86_n=200)  # different N -> mismatch
    from bench_core.compare_sweeps import _read_series_tables

    ok, warns = check_n_parity(_read_series_tables(manifest))
    assert ok is False
    assert any("not comparable" in w for w in warns)


# ---- end-to-end run -------------------------------------------------------


def test_run_comparison_writes_outputs(tmp_path):
    manifest = _two_series(tmp_path)
    out = tmp_path / "out"
    rc = run_comparison(manifest, out)
    assert rc == 0
    assert (out / "comparison-tidy.csv").exists()
    assert (out / "comparison-ratio-summary.csv").exists()
    assert (out / "comparison-trajectory-delta.csv").exists()
    assert (out / "comparison.xlsx").exists()
    # xlsx has the 4 sheets + LineChart on Per-ratio + ColorScale on heatmaps.
    from openpyxl import load_workbook

    wb = load_workbook(out / "comparison.xlsx")
    assert wb.sheetnames == ["Overview", "Per-ratio", "Component heatmaps", "Per-trajectory delta"]
    assert len(wb["Per-ratio"]._charts) >= 2  # absolute + degradation line charts
    assert len(wb["Component heatmaps"]._charts) == 0  # heatmaps are conditional fmt, not charts
    assert len(list(wb["Component heatmaps"].conditional_formatting)) >= 1


def test_run_comparison_refuses_mismatched_n(tmp_path):
    manifest = _two_series(tmp_path, x86_n=200)
    rc = run_comparison(manifest, tmp_path / "out")
    assert rc == 2  # refused
    # --allow-mismatched-n bypasses the gate (exposed via run_comparison kwarg).
    rc2 = run_comparison(manifest, tmp_path / "out2", allow_mismatched_n=True)
    assert rc2 == 0
    assert (tmp_path / "out2" / "comparison.xlsx").exists()


def test_main_shorthand_arm_x86(tmp_path):
    """--arm/--x86 shorthand synthesizes a 2-series manifest (arm baseline)."""
    _seed_series(tmp_path / "arm", ratios=[1], elapsed={(1, "t0"): 2.0})
    _seed_series(tmp_path / "x86", ratios=[1], elapsed={(1, "t0"): 1.0})
    out = tmp_path / "out"
    rc = main(["--arm", str(tmp_path / "arm"), "--x86", str(tmp_path / "x86"), "--output-dir", str(out)])
    assert rc == 0
    assert (out / "comparison.xlsx").exists()


def test_main_needs_manifest_or_arm_x86(tmp_path):
    assert main([]) == 2
