"""Unit tests for the per-trajectory degradation viewer.

Synthetic trajectory-detail.csv + ratio-summary.csv matching the oversub
driver contract (column names from ``bench_core.oversub.TRAJECTORY_COLUMNS``);
no live sweep runs. Drives ``run_traj_degradation`` end-to-end and
checks the xlsx + csv outputs.
"""
from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from bench_core.oversub import TRAJECTORY_COLUMNS
from bench_core.traj_degradation import (
    _aggregate,
    run_traj_degradation,
)


def _write_traj(path: Path, rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=TRAJECTORY_COLUMNS, extrasaction="ignore", restval="")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def _write_ratio_summary(path: Path, rows: list[dict]) -> None:
    cols = [
        "mode",
        "ratio",
        "attempted",
        "successful",
        "all_repeats_successful",
        "median_wall_sec",
        "median_tasks_per_sec",
        "median_peak_active",
        "median_queue_wait_sec",
        "time_degradation_vs_1_1_pct",
        "throughput_gain_vs_1_1_pct",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore", restval="")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def _traj_row(mode: str, ratio: int, traj_id: str, elapsed: float) -> dict:
    return {
        "mode": mode,
        "ratio": ratio,
        "repeat": 1,
        "trajectory_id": traj_id,
        "sandbox_index": 0,
        "round_id": 0,
        "n_steps": 3,
        "n_failed": 0,
        "n_timeout": 0,
        "success_rate": 1.0,
        "elapsed_sec": elapsed,
        "exec_sec": elapsed * 0.5,
        "resume_sec": elapsed * 0.1,
        "pause_sec": elapsed * 0.1,
        "slot_contention_wait_sec": elapsed * 0.1,
        "create_error_type": "",
        "kill_error_type": "",
    }


def _seed_sweep(root: Path, *, modes: list[str], ratios: list[int], tids: list[str], elapsed_fn, fleet_fn=None) -> None:
    """Write trajectory-detail.csv + optional ratio-summary.csv into root."""
    traj_rows = []
    for mode in modes:
        for ratio in ratios:
            for tid in tids:
                # 2 sandbox-runs per (mode,ratio,tid) so median is meaningful.
                for sb in (0, 1):
                    e = elapsed_fn(mode, ratio, tid)
                    if e is None:  # trajectory absent at this ratio (no baseline)
                        continue
                    r = _traj_row(mode, ratio, tid, e)
                    r["sandbox_index"] = sb
                    traj_rows.append(r)
    _write_traj(root / "trajectory-detail.csv", traj_rows)
    if fleet_fn is not None:
        rs_rows = []
        for mode in modes:
            for ratio in ratios:
                rs_rows.append(
                    {
                        "mode": mode,
                        "ratio": ratio,
                        "time_degradation_vs_1_1_pct": fleet_fn(mode, ratio),
                    }
                )
        _write_ratio_summary(root / "ratio-summary.csv", rs_rows)


def test_aggregate_degradation_and_blank_baseline(tmp_path):
    """A trajectory whose baseline_ratio is absent gets NaN degradation across
    ALL ratios (not 0%)."""
    _seed_sweep(
        tmp_path,
        modes=["lifecycle"],
        ratios=[1, 2, 3],
        tids=["tA", "tB"],
        # tA: 1.0 -> 1.2 -> 1.5 (degrades). tB: absent at ratio 1 (no baseline).
        elapsed_fn=lambda m, r, tid: {"tA": {1: 1.0, 2: 1.2, 3: 1.5}[r], "tB": {2: 2.0, 3: 3.0}.get(r)}[tid],
    )
    df = pd.read_csv(tmp_path / "trajectory-detail.csv")
    agg = _aggregate(df, baseline_ratio=1)
    a = agg[agg["trajectory_id"] == "tA"].set_index("ratio")
    # tA ratio 2: (1.2-1.0)/1.0*100 = 20.0
    assert round(float(a.loc[2, "degradation_pct"]), 1) == 20.0
    assert round(float(a.loc[3, "degradation_pct"]), 1) == 50.0
    b = agg[agg["trajectory_id"] == "tB"].set_index("ratio")
    # tB has no ratio-1 -> degradation NaN for ratio 2 AND 3; elapsed populated.
    assert all(pd.isna(b.loc[r, "degradation_pct"]) for r in [2, 3])
    assert float(b.loc[2, "elapsed_sec"]) == 2.0


def test_run_produces_xlsx_and_csv(tmp_path):
    _seed_sweep(
        tmp_path,
        modes=["lifecycle", "exec_only"],
        ratios=[1, 2, 3],
        tids=[f"t{i}" for i in range(6)],
        elapsed_fn=lambda m, r, tid: 1.0 * r * (1 + int(tid[1:]) * 0.1),
        fleet_fn=lambda m, r: (r - 1) * 15.0,  # fleet degrades 0/15/30 %
    )
    out = tmp_path / "out"
    rc = run_traj_degradation(tmp_path, out, top_n=4, heatmap_top_n=0)
    assert rc == 0
    assert (out / "traj-degradation.xlsx").exists()
    assert (out / "traj-degradation.csv").exists()

    wb = load_workbook(out / "traj-degradation.xlsx")
    # multi-mode -> each view suffixed with (mode).
    expected = ["Degradation (lifecycle)", "Degradation (exec_only)", "Top-N (lifecycle)", "Top-N (exec_only)"]
    for name in expected:
        assert name in wb.sheetnames, f"missing sheet {name}; got {wb.sheetnames}"
    # CSV: 2 modes x 6 trajectories x 3 ratios = 36 rows (per-run rows collapse to median).
    csv_df = pd.read_csv(out / "traj-degradation.csv")
    assert set(csv_df["mode"]) == {"lifecycle", "exec_only"}
    assert csv_df["trajectory_id"].nunique() == 6
    assert len(csv_df) == 36


def test_fleet_line_omitted_when_ratio_summary_absent(tmp_path):
    """No ratio-summary.csv -> Top-N sheet has no fleet_median column; rc still 0."""
    _seed_sweep(
        tmp_path,
        modes=["lifecycle"],
        ratios=[1, 2],
        tids=["tA", "tB"],
        elapsed_fn=lambda m, r, tid: {1: 1.0, 2: 1.3}[r],
        fleet_fn=None,  # don't write ratio-summary.csv
    )
    out = tmp_path / "out"
    rc = run_traj_degradation(tmp_path, out, top_n=2)
    assert rc == 0
    wb = load_workbook(out / "traj-degradation.xlsx")
    ws = wb["Top-N"]
    header = [c.value for c in ws[1]]
    assert "fleet_median" not in header
    assert header[0] == "ratio"


def test_heatmap_top_n_caps_rows(tmp_path):
    """--heatmap-top-n 2 limits Degradation heatmap to 2 worst trajectories."""
    _seed_sweep(
        tmp_path,
        modes=["lifecycle"],
        ratios=[1, 2, 3],
        tids=[f"t{i}" for i in range(6)],
        # t5 degrades most (highest elapsed growth), should be row 1.
        elapsed_fn=lambda m, r, tid: {1: 1.0, 2: 1.0 + 0.1 * int(tid[1:]), 3: 1.0 + 0.3 * int(tid[1:])}[r],
    )
    out = tmp_path / "out"
    rc = run_traj_degradation(tmp_path, out, top_n=3, heatmap_top_n=2)
    assert rc == 0
    wb = load_workbook(out / "traj-degradation.xlsx")
    ws = wb["Degradation"]
    # header row + 2 data rows = 3 rows total.
    assert ws.max_row == 3
    # Worst trajectory (t5) first.
    assert ws.cell(row=2, column=1).value == "t5"


def test_missing_sweep_dir_returns_2(tmp_path):
    assert run_traj_degradation(tmp_path / "nope", tmp_path / "out") == 2
