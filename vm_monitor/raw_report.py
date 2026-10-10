"""raw_data/*.csv -> derived-rate DataFrames (spotbox-style host_resources consumer).

Reads the always-on raw /proc + per-VM + per-disk CSV dumps and derives rates
via delta (reusing base._compute_pressure_rates + base._compute_disk_io_rates).
Pure file input -- does not depend on in-memory history, post-hoc rerunnable.
The xlsx assembly (build_host_resources_xlsx: 3 spotbox-style sheets with
openpyxl LineCharts) reuses the derived DataFrames above. All derivation is
delta-based on persistent cumulative counters -- pure file input, post-hoc
rerunnable, independent of in-memory history and of resource_report.xlsx.
"""
from __future__ import annotations

import logging
import math
from pathlib import Path

import pandas as pd

from vm_monitor.base import (
    _BYTES_PER_MIB,
    _CLK_TCK,
    _DISK_FIELDS,
    _compute_disk_io_rates,
    _compute_pressure_rates,
)

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Calibration knobs (ponytail: physical world needs tuning, not hard-coded).
MAX_INTERVAL_S = 10.0  # gap guard: per-key delta over a gap longer than this -> NaN
MAX_VM_SERIES = 8  # per-VM chart series cap (Excel render perf); used in Task 3+

# /sys/block/<dev>/stat fields that are cumulative monotonic counters.
# `inflight` (current I/Os in progress) is a GAUGE -- normal 3->1 fluctuation
# must NOT trip the reset guard, else busy-disk intervals get NaN-holed.
_DISK_GUARD_FIELDS = tuple(f for f in _DISK_FIELDS if f != "inflight")

_MEMINFO_KB_PER_GIB = 2**20  # /proc/meminfo fields are in kB; /2**20 -> GiB

_VMSTAT_FIELDS = [
    "pswpin",
    "pswpout",
    "pgscan_kswapd",
    "pgscan_direct",
    "pgscan_direct_throttle",
    "pgsteal_kswapd",
    "pgsteal_direct",
    "workingset_refault_file",
]


def _f(d: dict, k: str) -> float:
    """Best-effort float of a CSV cell (NaN if empty/non-numeric)."""
    v = d.get(k)
    try:
        return float(v)
    except (TypeError, ValueError):
        return math.nan


def _guarded(
    cur: dict, prev: dict, dt: float, key: str, fields: list[str], max_interval: float = MAX_INTERVAL_S
) -> bool:
    """True if this delta must be NaNed: dt<=0, dt>max_interval (gap), or any
    cumulative field cur<prev (counter reset -- e.g. VM pid restart zeroes
    utime/stime, device remount zeroes disk stat). Logs a WARNING per case."""
    if dt <= 0:
        logger.warning("counter reset/gap for %s (dt<=0)", key)
        return True
    if dt > max_interval:
        logger.warning("gap detected for %s (dt=%.1fs > %.1fs)", key, dt, max_interval)
        return True
    for f in fields:
        if _f(cur, f) < _f(prev, f):
            logger.warning("counter reset for %s at %s (cur<prev in %s)", key, cur.get("timestamp"), f)
            return True
    return False


def _read_csv(raw_dir: Path, fname: str) -> pd.DataFrame | None:
    p = raw_dir / fname
    if not p.exists() or p.stat().st_size == 0:
        logger.warning("raw CSV missing/empty: %s -> skipping derived sheet", p)
        return None
    try:
        return pd.read_csv(p)
    except (pd.errors.ParserError, OSError, ValueError) as e:
        logger.warning("raw CSV unparseable: %s (%s) -> skipping", p, e)
        return None


def _aggregate_cores(ublk_df: pd.DataFrame, cur_ts: float, dt: float) -> float:
    """Sum ublk-daemon cores at sample boundary cur_ts (ublk is 1s sub-sampled;
    take the last row <= cur_ts per pid, delta vs its predecessor)."""
    if ublk_df is None or ublk_df.empty or not dt:
        return math.nan
    # ponytail: re-sorts the whole ublk_df per call (O(n log n) x n_samples);
    # acceptable for few-daemon hosts at ~1Hz over minutes-to-hours. Hoist the
    # sort out of the loop (pre-sort once, groupby preserves order) if a long
    # high-frequency run makes this hot.
    ublk_df = ublk_df.sort_values("timestamp")
    ublk_df["timestamp"] = ublk_df["timestamp"].astype(float)
    total = 0.0
    nans = 0
    for pid, g in ublk_df.groupby("pid"):
        g = g[g["timestamp"] <= cur_ts]
        if len(g) < 2:
            nans += 1
            continue
        prev, cur = g.iloc[-2], g.iloc[-1]
        d = (
            _f(cur.to_dict(), "utime")
            + _f(cur.to_dict(), "stime")
            - _f(prev.to_dict(), "utime")
            - _f(prev.to_dict(), "stime")
        )
        sub_dt = cur["timestamp"] - prev["timestamp"]
        if _guarded(cur.to_dict(), prev.to_dict(), sub_dt, f"ublk pid={pid}", ["utime", "stime"]):
            nans += 1
            continue
        total += d / _CLK_TCK / sub_dt
    return total if nans == 0 else math.nan


def _derive_host(raw_dir: Path) -> pd.DataFrame:
    """Derive host-wide rates from host_mem_raw + host_cpu_raw + host_vmstat_raw
    + ublk_cpu_raw (aggregate). One row per timestamp (sorted). First row NaN."""
    cpu = _read_csv(raw_dir, "host_cpu_raw.csv")
    mem = _read_csv(raw_dir, "host_mem_raw.csv")
    vmstat = _read_csv(raw_dir, "host_vmstat_raw.csv")
    ublk = _read_csv(raw_dir, "ublk_cpu_raw.csv")
    if cpu is None or cpu.empty:
        return pd.DataFrame()
    cpu = cpu.sort_values("timestamp").reset_index(drop=True)
    ts = cpu["timestamp"].astype(float).to_numpy()
    out = {
        c: [math.nan]
        for c in (
            "cpu_busy_pct",
            "cpu_iowait_pct",
            "mem_used_gb",
            "mem_cache_gb",
            "buffers_mb",
            "dirty_mb",
            "swap_in_mib_s",
            "swap_out_mib_s",
            "page_scan_mib_s",
            "page_reclaim_mib_s",
            "file_refault_mib_s",
            "ublk_cores",
            "vm_cores",
        )
    }

    _CPU8 = ["cpu_user", "cpu_nice", "cpu_system", "cpu_idle", "cpu_iowait", "cpu_irq", "cpu_softirq", "cpu_steal"]
    for i in range(1, len(cpu)):
        dt = ts[i] - ts[i - 1]
        cur_c, prev_c = cpu.iloc[i].to_dict(), cpu.iloc[i - 1].to_dict()
        if not _guarded(cur_c, prev_c, dt, "host_cpu", _CPU8):
            d_total = sum(_f(cur_c, f) - _f(prev_c, f) for f in _CPU8)
            d_idle_iowait = (
                _f(cur_c, "cpu_idle") + _f(cur_c, "cpu_iowait") - _f(prev_c, "cpu_idle") - _f(prev_c, "cpu_iowait")
            )
            out["cpu_busy_pct"].append(100.0 * (1 - d_idle_iowait / d_total) if d_total else math.nan)
            out["cpu_iowait_pct"].append(
                100.0 * (_f(cur_c, "cpu_iowait") - _f(prev_c, "cpu_iowait")) / d_total if d_total else math.nan
            )
        else:
            out["cpu_busy_pct"].append(math.nan)
            out["cpu_iowait_pct"].append(math.nan)
        # mem is a gauge (no delta); aligned by row index to cpu since host_*_raw
        # are written together per-sample in one enqueue_proc call. ponytail:
        # positional alignment assumes no per-/proc read failure drops a row;
        # switch to pd.merge_asof on timestamp if that ever happens.
        if mem is not None and i < len(mem):
            md = mem.iloc[i].to_dict()
            out["mem_used_gb"].append((_f(md, "MemTotal") - _f(md, "MemAvailable")) / _MEMINFO_KB_PER_GIB)
            out["mem_cache_gb"].append((_f(md, "Cached") + _f(md, "SReclaimable")) / _MEMINFO_KB_PER_GIB)
            out["buffers_mb"].append(_f(md, "Buffers") / 1024)
            out["dirty_mb"].append(_f(md, "Dirty") / 1024)
        else:
            for k in ("mem_used_gb", "mem_cache_gb", "buffers_mb", "dirty_mb"):
                out[k].append(math.nan)
        if vmstat is not None and i < len(vmstat):
            cur_v, prev_v = vmstat.iloc[i].to_dict(), vmstat.iloc[i - 1].to_dict()
            if _guarded(cur_v, prev_v, dt, "host_vmstat", _VMSTAT_FIELDS):
                for k in (
                    "swap_in_mib_s",
                    "swap_out_mib_s",
                    "page_scan_mib_s",
                    "page_reclaim_mib_s",
                    "file_refault_mib_s",
                ):
                    out[k].append(math.nan)
            else:
                d_pswpin = _f(cur_v, "pswpin") - _f(prev_v, "pswpin")
                d_pswpout = _f(cur_v, "pswpout") - _f(prev_v, "pswpout")
                out["swap_in_mib_s"].append(max(0, d_pswpin) * 4096 / _BYTES_PER_MIB / dt if dt else math.nan)
                out["swap_out_mib_s"].append(max(0, d_pswpout) * 4096 / _BYTES_PER_MIB / dt if dt else math.nan)
                pr = _compute_pressure_rates(cur_v, prev_v, dt)
                out["page_scan_mib_s"].append(pr["page_scan_mib_s"])
                out["page_reclaim_mib_s"].append(pr["page_reclaim_mib_s"])
                out["file_refault_mib_s"].append(pr["file_refault_mib_s"])
        else:
            for k in ("swap_in_mib_s", "swap_out_mib_s", "page_scan_mib_s", "page_reclaim_mib_s", "file_refault_mib_s"):
                out[k].append(math.nan)
        out["ublk_cores"].append(_aggregate_cores(ublk, ts[i], dt) if ublk is not None else math.nan)
        out["vm_cores"].append(math.nan)  # filled by per-VM aggregate in the orchestrator (Task 5)
    df = pd.DataFrame(out)
    df.insert(0, "t_s", [0.0] + [ts[i] - ts[0] for i in range(1, len(ts))])
    df.insert(1, "timestamp", cpu["timestamp"].tolist())
    return df


def _derive_per_vm(raw_dir: Path, max_series: int = MAX_VM_SERIES) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Per-VM CPU cores (delta of utime+stime /_CLK_TCK/dt, guarded) + PSS GB
    (gauge, plotted directly -- no delta). long -> wide pivot (rows indexed by
    timestamp, cols=vm_name). VM count > max_series -> top-K by CPU busy mean
    (descending) + 'other' = sum(skipna, min_count=1) of the rest. Both the CPU
    and PSS charts share this VM set (ranked by CPU busy, NOT PSS) so the two
    side-by-side charts show the same VMs -- ranking each by its own metric
    would misalign the sets and confuse the reader."""
    vm = _read_csv(raw_dir, "vm_cpu_raw.csv")
    if vm is None or vm.empty:
        return pd.DataFrame(), pd.DataFrame()
    vm = vm.sort_values(["vm_name", "timestamp"]).copy()
    vm["timestamp"] = vm["timestamp"].astype(float)
    vm["cores"] = math.nan
    vm["pss_gb"] = math.nan
    for name, g in vm.groupby("vm_name"):
        g = g.sort_values("timestamp")
        prev = None
        c_col, p_col = [], []
        for _, row in g.iterrows():
            rd = row.to_dict()
            if prev is None:
                c_col.append(math.nan)
            else:
                dt = float(row["timestamp"]) - float(prev["timestamp"])
                if _guarded(rd, prev, dt, f"vm={name}", ["utime", "stime"]):
                    c_col.append(math.nan)
                else:
                    d = _f(rd, "utime") + _f(rd, "stime") - _f(prev, "utime") - _f(prev, "stime")
                    c_col.append(d / _CLK_TCK / dt if dt else math.nan)
            p_col.append(_f(rd, "pss_mb") / 1024)
            prev = rd
        vm.loc[g.index, "cores"] = c_col
        vm.loc[g.index, "pss_gb"] = p_col

    p_cores = vm.pivot_table(index="timestamp", columns="vm_name", values="cores", aggfunc="first", dropna=False)
    p_pss = vm.pivot_table(index="timestamp", columns="vm_name", values="pss_gb", aggfunc="first", dropna=False)
    # rank VMs by mean CPU busy (cores) -- shared set for both charts
    means = p_cores.mean().sort_values(ascending=False)
    top = list(means.head(max_series).index)
    rest = [v for v in p_cores.columns if v not in top]
    out_cores = p_cores[top].copy()
    out_pss = p_pss[top].copy()
    if rest:
        out_cores["other"] = p_cores[rest].sum(axis=1, skipna=True, min_count=1)
        out_pss["other"] = p_pss[rest].sum(axis=1, skipna=True, min_count=1)
    out_cores = out_cores.sort_index()
    out_pss = out_pss.sort_index()
    t0 = out_cores.index.min()
    out_cores.insert(0, "t_s", out_cores.index - t0)
    out_pss.insert(0, "t_s", out_pss.index - t0)
    return out_cores, out_pss


def _derive_disk(raw_dir: Path) -> pd.DataFrame:
    """Per-device r/w MB/s + util% via base._compute_disk_io_rates. long CSV ->
    {ts: {dev: {field}}}; consecutive ts snapshots -> rates, with a per-device
    counter-reset guard (cur<prev -> NaN for that dev that interval, no
    negative spike on device remount)."""
    disk = _read_csv(raw_dir, "disk_io_raw.csv")
    if disk is None or disk.empty:
        return pd.DataFrame()
    disk["timestamp"] = disk["timestamp"].astype(float)
    # build per-timestamp snapshot {dev: {field: val}}
    snaps = {}
    for ts, g in disk.groupby("timestamp"):
        snaps[ts] = {r["device"]: {f: r[f] for f in _DISK_FIELDS} for _, r in g.iterrows()}
    ts_sorted = sorted(snaps)
    devs = sorted({dev for s in snaps.values() for dev in s})
    cols = {}
    for dev in devs:
        cols[f"{dev}_r_mb_s"] = [math.nan]
        cols[f"{dev}_w_mb_s"] = [math.nan]
        cols[f"{dev}_util_pct"] = [math.nan]
    for i in range(1, len(ts_sorted)):
        cur_ts, prev_ts = ts_sorted[i], ts_sorted[i - 1]
        dt = cur_ts - prev_ts
        for dev in devs:
            for sfx in ("r_mb_s", "w_mb_s", "util_pct"):
                cols[f"{dev}_{sfx}"].append(math.nan)
        cur_snap, prev_snap = snaps[cur_ts], snaps[prev_ts]
        # per-dev reset guard, then reuse base compute (single-dev dict).
        # Guard over cumulative counters ONLY (_DISK_GUARD_FIELDS) -- inflight is
        # a gauge and is read as such by _compute_disk_io_rates, so a normal drop
        # must not NaN-hole the interval. Inject timestamp for the WARNING log.
        for dev in devs:
            if dev not in cur_snap or dev not in prev_snap:
                continue
            if _guarded(
                {"timestamp": cur_ts, **cur_snap[dev]},
                {"timestamp": prev_ts, **prev_snap[dev]},
                dt,
                f"disk={dev}",
                _DISK_GUARD_FIELDS,
            ):
                continue
            rates = _compute_disk_io_rates({dev: cur_snap[dev]}, {dev: prev_snap[dev]}, dt)[dev]
            cols[f"{dev}_r_mb_s"][-1] = rates["r_mb_s"]
            cols[f"{dev}_w_mb_s"][-1] = rates["w_mb_s"]
            cols[f"{dev}_util_pct"][-1] = rates["util_pct"]
    df = pd.DataFrame(cols)
    df.insert(0, "t_s", [0.0] + [ts_sorted[i] - ts_sorted[0] for i in range(1, len(ts_sorted))])
    return df


_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="F9FAFB")


def _put_row(ws, row, values, header=False):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        if header:
            c.fill = _HEADER_FILL
            c.font = _HEADER_FONT
    return row + 1


def _write_table(ws, df, headers):
    """Write df rows under a styled header at row 1; return last data-row index."""
    r = _put_row(ws, 1, headers, header=True)
    for _, row in df.iterrows():
        r = _put_row(ws, r, [row.get(h) for h in headers])
    ws.freeze_panes = "A2"
    return r - 1  # last data row


def _add_line_chart(ws, title, y_title, cat_col, data_cols, header_row, n_data_rows, anchor):
    """openpyxl LineChart bound to worksheet cells (spotbox style). ``header_row``
    carries the series names (titles_from_data); data occupies header_row+1 ..
    header_row+n_data_rows. X axis = cat_col (t (s)). Stacked vertically via anchor."""
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = y_title
    ch.x_axis.title = "t (s)"
    ch.height = 8
    ch.width = 16
    last = header_row + n_data_rows
    for col in data_cols:
        ch.add_data(Reference(ws, min_col=col, min_row=header_row, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=header_row + 1, max_row=last))
    ws.add_chart(ch, anchor)


def _build_host_sheet(wb, df):
    if df.empty:
        return
    ws = wb.create_sheet("Host resources")
    headers = [
        "t_s",
        "timestamp",
        "cpu_busy_pct",
        "cpu_iowait_pct",
        "mem_used_gb",
        "mem_cache_gb",
        "buffers_mb",
        "dirty_mb",
        "swap_in_mib_s",
        "swap_out_mib_s",
        "page_scan_mib_s",
        "page_reclaim_mib_s",
        "file_refault_mib_s",
        "ublk_cores",
        "vm_cores",
    ]
    _write_table(ws, df, headers)
    anchor = get_column_letter(len(headers) + 2)
    n = len(df)
    _add_line_chart(ws, "Host memory", "GB/MB", 1, (5, 6, 7, 8), 1, n, f"{anchor}2")
    _add_line_chart(ws, "Host CPU", "%", 1, (3, 4), 1, n, f"{anchor}20")
    _add_line_chart(ws, "Page-cache pressure", "MiB/s", 1, (11, 12, 13), 1, n, f"{anchor}38")
    _add_line_chart(ws, "Swap in/out", "MiB/s", 1, (9, 10), 1, n, f"{anchor}56")
    _add_line_chart(ws, "CPU cores (ublk + VM)", "cores", 1, (14, 15), 1, n, f"{anchor}74")


def _build_per_vm_sheet(wb, cores_df, pss_df):
    if cores_df is None or cores_df.empty:
        return
    ws = wb.create_sheet("Per-VM")
    cores_headers = ["t_s", *cores_df.columns[1:]]  # t_s + VM names + 'other'
    last_core = _write_table(ws, cores_df, cores_headers)
    anchor = get_column_letter(len(cores_headers) + 2)
    # per-vm table has NO timestamp column -> first VM is col 2 (not 3).
    cpu_cols = tuple(range(2, len(cores_headers) + 1))
    _add_line_chart(ws, "Per-VM CPU cores", "cores", 1, cpu_cols, 1, len(cores_df), f"{anchor}2")
    # PSS table written BELOW the cores table; its chart must reference the PSS
    # rows (header_row + first_data_row), not the cores rows at the top.
    pss_headers = ["t_s", *pss_df.columns[1:]]
    pss_hdr = last_core + 3
    r = _put_row(ws, pss_hdr, pss_headers, header=True)
    for _, row in pss_df.iterrows():
        r = _put_row(ws, r, [row.get(h) for h in pss_headers])
    pss_cols = tuple(range(2, len(pss_headers) + 1))
    # Floor the PSS chart anchor at row 20 so it clears the 8cm cores chart
    # (anchored at row 2, spans ~rows 2-17) when the cores table is short.
    # ponytail: production has hundreds of samples so pss_hdr+1 >> 20; the floor
    # only matters for tiny test fixtures.
    pss_anchor_row = max(pss_hdr + 1, 20)
    _add_line_chart(ws, "Per-VM PSS", "GB", 1, pss_cols, pss_hdr, len(pss_df), f"{anchor}{pss_anchor_row}")


def _build_disk_sheet(wb, df):
    if df.empty:
        return
    ws = wb.create_sheet("Disk IO")
    headers = list(df.columns)
    _write_table(ws, df, headers)
    anchor = get_column_letter(len(headers) + 2)
    n = len(df)
    rw_cols = tuple(i for i, h in enumerate(headers, 1) if h.endswith("_r_mb_s") or h.endswith("_w_mb_s"))
    util_cols = tuple(i for i, h in enumerate(headers, 1) if h.endswith("_util_pct"))
    if rw_cols:
        _add_line_chart(ws, "Disk r/w MB/s", "MB/s", 1, rw_cols, 1, n, f"{anchor}2")
    if util_cols:
        _add_line_chart(ws, "Disk util %", "%", 1, util_cols, 1, n, f"{anchor}20")


def build_host_resources_xlsx(raw_dir, out_path) -> Path:
    """Read raw_data/*.csv -> derive -> write host_resources.xlsx (3 spotbox-style
    sheets). Independent of resource_report.xlsx (separate filename preserves
    bench-core's reap signal) and of in-memory history (post-hoc rerunnable).
    Degrades: a missing CSV skips its sheet (WARNING), never raises."""
    raw_dir = Path(raw_dir)
    out_path = Path(out_path)
    host_df = _derive_host(raw_dir)
    cores_df, pss_df = pd.DataFrame(), pd.DataFrame()
    try:
        cores_df, pss_df = _derive_per_vm(raw_dir)
    except (OSError, ValueError, KeyError, pd.errors.ParserError) as e:
        logger.warning("per-VM derivation failed for raw_dir=%s: %s", raw_dir, e)
        cores_df, pss_df = pd.DataFrame(), pd.DataFrame()
    disk_df = _derive_disk(raw_dir)
    # Fill host_df.vm_cores from the per-VM cores aggregate (sum across VMs per
    # timestamp). host_df is RangeIndex; cores_df is timestamp-indexed -> align
    # by position (same sample cadence). ponytail: positional fill assumes host
    # and vm CSVs share sample cadence (true: both written by collect_sample);
    # a cadence mismatch leaves vm_cores NaN rather than misaligning silently.
    if not cores_df.empty and not host_df.empty:
        vm_agg = (
            cores_df.drop(columns=["t_s"], errors="ignore")
            .select_dtypes("number")
            .sum(axis=1, skipna=True, min_count=1)
        )
        if len(vm_agg) == len(host_df):
            host_df["vm_cores"] = vm_agg.to_numpy()
    wb = Workbook()
    default = wb.active
    _build_host_sheet(wb, host_df)
    _build_per_vm_sheet(wb, cores_df, pss_df)
    _build_disk_sheet(wb, disk_df)
    # If no sheet was built (all raw CSVs missing/empty), there is nothing to
    # write -- log + return without saving a misleading empty workbook. Also
    # avoids wb.save() IndexError on a 0-sheet book (default was not removed).
    if len(wb.worksheets) == 1:  # only the default -> no sheet built at all
        logger.warning("no raw CSV data found in raw_dir=%s; host_resources.xlsx not written", raw_dir)
        return out_path
    wb.remove(default)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logger.info("host_resources.xlsx written: %s", out_path)
    return out_path
