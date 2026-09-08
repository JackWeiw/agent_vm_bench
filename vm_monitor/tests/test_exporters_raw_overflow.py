"""Row-overflow guard for the ``Raw_VM_Data`` sheet.

``Raw_VM_Data`` carries one row per (sample, VM), so a large fleet
(384 VMs * 14k samples ~= 5.4M rows) exceeds Excel's 1,048,576-row cap.
Without the guard, ``to_excel`` raises inside the shared ``ExcelWriter``
block and aborts the ENTIRE workbook -- every other sheet + charts lost
because ONE sheet was too big. The guard skips just that sheet and points
at the raw CSV (same schema, no row limit) so the rest of the workbook
survives.
"""

import os
import tempfile
import unittest
from unittest import mock

from vm_monitor.base import VMMonitorBase
from vm_monitor.exporters import PANDAS_AVAILABLE, _EXCEL_MAX_ROWS, export_to_excel

try:
    import pandas as pd
except ImportError:  # pragma: no cover
    pd = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


class DummyMonitor(VMMonitorBase):
    """Concrete subclass for testing (VMMonitorBase is abstract)."""

    def get_vms_realtime(self):
        return []

    def get_process_names(self):
        return ("test_process",)

    def extract_vm_id(self, pid, cmdline):
        return "vm0"

    def get_monitor_title(self):
        return "DummyMonitor"

    def get_no_vm_message(self):
        return "No VMs detected"

    def get_csv_filename_prefix(self):
        return "dummy_monitor"


def _record(i):
    """One per-VM-per-sample record (minimal but schema-valid)."""
    return {
        "timestamp": f"2026-09-07 10:00:{i:02d}",
        "vm_name": "vm0",
        "pid": 100,
        "cpu_percent": 10.0,
        "memory_mb": 2048.0,
        "memory_huge_mb": 1024.0,
        "memory_private_mb": 1000.0,
        "memory_heap_mb": 50.0,
        "memory_swapcache_mb": 10.0,
        "memory_per_numa": {0: {"total_mb": 1000.0}},
        "memory_swapcache_per_numa": {0: 4.0},
        "status": "running",
    }


def _mem_record(i):
    """One host_mem_detail_history sample (minimal but schema-valid)."""
    return {
        "ts": f"2026-09-07 10:00:{i:02d}",
        "cached_mb": 1000.0 + i,
        "buffers_mb": 100.0,
        "dirty_mb": 10.0,
        "writeback_mb": 1.0,
    }


@unittest.skipUnless(PANDAS_AVAILABLE and load_workbook is not None, "pandas/openpyxl required")
class TestRawVmDataOverflow(unittest.TestCase):
    """The overflow guard skips Raw_VM_Data but keeps every other sheet."""

    def setUp(self):
        self.monitor = DummyMonitor()
        self.log_dir = tempfile.mkdtemp(prefix="vm_monitor_overflow_")
        self.output_file = os.path.join(self.log_dir, "resource_report.xlsx")

    def tearDown(self):
        for f in os.listdir(self.log_dir):
            try:
                os.unlink(os.path.join(self.log_dir, f))
            except PermissionError:
                pass
        os.rmdir(self.log_dir)

    def _export(self):
        with mock.patch("vm_monitor.exporters.parse_all_logs", return_value={}):
            return export_to_excel(self.monitor, self.log_dir, numa_nodes=[0], output_file=self.output_file)

    def test_oversized_raw_vm_data_skipped_other_sheets_survive(self):
        """When monitor.data exceeds the row cap, Raw_VM_Data is skipped (with
        a warning pointing at the raw CSV) but the rest of the workbook +
        charts are written intact. Previously this raised inside the
        ExcelWriter block and lost the ENTIRE workbook."""
        # Lower the cap so we can trigger the guard without building >1M
        # records. 5 records, cap at 4 -> guard fires.
        self.monitor.data = [_record(i) for i in range(5)]
        with mock.patch("vm_monitor.exporters._EXCEL_MAX_ROWS", 4), mock.patch("builtins.print") as mock_print:
            self._export()

        # The skip warning names the sheet, the row count, and the CSV fallback.
        printed = " ".join(str(c[0][0]) for c in mock_print.call_args_list)
        self.assertIn("Raw_VM_Data", printed)
        self.assertIn("5", printed)  # actual row count
        self.assertIn("CSV", printed)
        self.assertIn(self.log_dir, printed)

        # Raw_VM_Data must be absent; other sheets must survive.
        with pd.ExcelFile(self.output_file) as xf:
            names = xf.sheet_names
        self.assertNotIn("Raw_VM_Data", names)
        self.assertIn("Summary", names)
        self.assertIn("VM_Stats", names)

    def test_under_cap_raw_vm_data_written_normally(self):
        """At-or-below the row cap the sheet is written as before -- the guard
        does not false-positive and the sheet is present."""
        self.monitor.data = [_record(i) for i in range(4)]
        with mock.patch("vm_monitor.exporters._EXCEL_MAX_ROWS", 4):
            self._export()
        with pd.ExcelFile(self.output_file) as xf:
            names = xf.sheet_names
        self.assertIn("Raw_VM_Data", names)

    def test_real_cap_constant_matches_excel_limit(self):
        """The shipped constant must equal Excel's hard row limit (1,048,576).
        Guard against an accidental bump that would let oversized sheets
        through to openpyxl's hard refusal."""
        self.assertEqual(_EXCEL_MAX_ROWS, 1_048_576)


@unittest.skipUnless(PANDAS_AVAILABLE and load_workbook is not None, "pandas/openpyxl required")
class TestTimelineSheetOverflow(unittest.TestCase):
    """The per-sample timeline sheets (memory/disk/pressure/swap) also carry one
    row per sample, so a very long collection can exceed the row cap. The guard
    must skip just that sheet, write the full data to a CSV fallback (no row
    limit), and let the rest of the workbook + charts survive -- the same
    contract as Raw_VM_Data, but the timeline sheets have no pre-existing raw
    CSV, so the guard writes one itself.
    """

    def setUp(self):
        self.monitor = DummyMonitor()
        self.log_dir = tempfile.mkdtemp(prefix="vm_monitor_timeline_")
        self.output_file = os.path.join(self.log_dir, "resource_report.xlsx")

    def tearDown(self):
        for f in os.listdir(self.log_dir):
            try:
                os.unlink(os.path.join(self.log_dir, f))
            except PermissionError:
                pass
        os.rmdir(self.log_dir)

    def _export(self):
        with mock.patch("vm_monitor.exporters.parse_all_logs", return_value={}):
            return export_to_excel(self.monitor, self.log_dir, numa_nodes=[0], output_file=self.output_file)

    def test_oversized_timeline_sheet_skipped_csv_written(self):
        """When a per-sample timeline exceeds the row cap, the sheet is skipped
        (with a warning naming the row count + the CSV fallback path) but the
        full data lands in ``<log_dir>/<sheet>.csv`` and the rest of the
        workbook + charts are written intact. Previously this raised inside
        the ExcelWriter block and lost the ENTIRE workbook."""
        # Lower the cap so we can trigger the guard without building >1M
        # records. 5 records, cap at 4 -> guard fires.
        self.monitor.host_mem_detail_history = [_mem_record(i) for i in range(5)]
        # A couple of VM records so VM_Stats (an aggregated sheet) is also
        # written -- lets us assert "the rest of the workbook survived".
        self.monitor.data = [_record(i) for i in range(2)]
        with mock.patch("vm_monitor.exporters._EXCEL_MAX_ROWS", 4), mock.patch("builtins.print") as mock_print:
            self._export()

        printed = " ".join(str(c[0][0]) for c in mock_print.call_args_list)
        self.assertIn("Host_Mem_Timeline", printed)
        self.assertIn("5", printed)  # actual row count
        self.assertIn("CSV", printed)

        # The full data was written to a CSV fallback (same schema, no row limit).
        csv_path = os.path.join(self.log_dir, "Host_Mem_Timeline.csv")
        self.assertTrue(os.path.exists(csv_path), f"expected CSV fallback at {csv_path}")
        import csv as _csv

        with open(csv_path, encoding="utf-8") as fh:
            rows = list(_csv.reader(fh))
        self.assertEqual(len(rows), 6)  # header + 5 data rows
        self.assertEqual(rows[0][0], "Timestamp")
        self.assertEqual(rows[1][0], "2026-09-07 10:00:00")

        # Host_Mem_Timeline must be absent; other sheets must survive.
        with pd.ExcelFile(self.output_file) as xf:
            names = xf.sheet_names
        self.assertNotIn("Host_Mem_Timeline", names)
        self.assertIn("Summary", names)
        self.assertIn("VM_Stats", names)

    def test_under_cap_timeline_sheet_written_normally(self):
        """At-or-below the row cap the timeline sheet is written as before --
        the guard does not false-positive and no CSV fallback is created."""
        self.monitor.host_mem_detail_history = [_mem_record(i) for i in range(4)]
        with mock.patch("vm_monitor.exporters._EXCEL_MAX_ROWS", 4):
            self._export()
        with pd.ExcelFile(self.output_file) as xf:
            names = xf.sheet_names
        self.assertIn("Host_Mem_Timeline", names)
        # No fallback CSV when the sheet fit.
        self.assertFalse(os.path.exists(os.path.join(self.log_dir, "Host_Mem_Timeline.csv")))


if __name__ == "__main__":
    unittest.main()
