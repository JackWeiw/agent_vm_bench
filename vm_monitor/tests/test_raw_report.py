"""Tests for raw_report: raw_data CSV -> derived-rate xlsx (spotbox-style)."""
from __future__ import annotations

import csv
import logging
import math
import openpyxl
import os

from vm_monitor.raw_report import (
    _derive_disk,
    _derive_host,
    _derive_per_vm,
    _guarded,
    build_host_resources_xlsx,
)


def _write(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def test_guarded_flags_gap_and_reset(caplog):
    """_guarded returns True (-> NaN) on dt>max_interval, dt<=0, or cur<prev."""
    with caplog.at_level(logging.WARNING):
        # gap: dt=20 > max_interval=10
        assert _guarded({"u": 200}, {"u": 100}, 20.0, "vm=x", ["u"]) is True
        # reset: cur<prev (dt ok)
        assert _guarded({"u": 50}, {"u": 100}, 2.0, "vm=x", ["u"]) is True
        # dt<=0
        assert _guarded({"u": 110}, {"u": 100}, 0.0, "vm=x", ["u"]) is True
        # clean delta -> False
        assert _guarded({"u": 110}, {"u": 100}, 2.0, "vm=x", ["u"]) is False
    assert any("gap detected" in r for r in caplog.messages)
    assert any("counter reset" in r for r in caplog.messages)


def test_derive_host_cpu_mem_swap(tmp_path):
    """host CPU% / mem GB / swap MiB/s / pressure / ublk cores from synthetic
    cumulative CSVs. Counters are MONOTONIC so the reset guard does not fire."""
    d = tmp_path
    # Monotonic counters: idle 8000->8010, iowait 100->102, sum8 8285->8326 (d_total=41)
    _write(
        d / "host_cpu_raw.csv",
        [
            "timestamp",
            "cpu_user",
            "cpu_nice",
            "cpu_system",
            "cpu_idle",
            "cpu_iowait",
            "cpu_irq",
            "cpu_softirq",
            "cpu_steal",
            "cpu_guest",
            "cpu_guest_nice",
            "ctxt",
            "btime",
            "processes",
            "procs_running",
            "procs_blocked",
            "softirq_total",
        ],
        [
            ["0", "100", "0", "50", "8000", "100", "10", "20", "5", "0", "0", "1", "0", "0", "0", "0", "0"],
            ["2", "120", "0", "55", "8010", "102", "11", "22", "6", "0", "0", "2", "0", "0", "0", "0", "0"],
        ],
    )
    _write(
        d / "host_mem_raw.csv",
        ["timestamp", "MemTotal", "MemAvailable", "MemFree", "Cached", "SReclaimable", "Buffers", "Dirty", "Writeback"],
        [
            ["0", "2097152", "1048576", "512000", "200000", "50000", "30000", "1000", "0"],
            ["2", "2097152", "943718", "450000", "210000", "52000", "32000", "2000", "0"],
        ],
    )
    _write(
        d / "host_vmstat_raw.csv",
        [
            "timestamp",
            "pswpin",
            "pswpout",
            "pgscan_kswapd",
            "pgscan_direct",
            "pgscan_direct_throttle",
            "pgsteal_kswapd",
            "pgsteal_direct",
            "workingset_refault_file",
        ],
        [["0", "100", "0", "0", "0", "0", "0", "0", "0"], ["2", "300", "0", "2048", "0", "0", "0", "0", "0"]],
    )
    _write(
        d / "ublk_cpu_raw.csv",
        ["timestamp", "pid", "utime", "stime"],
        [["0", "4242", "0", "0"], ["2", "4242", "100", "100"]],
    )

    df = _derive_host(d)
    r = df.iloc[1]
    # d(idle+iowait)=12, d_total8=41 -> busy% = 100*(1 - 12/41)
    assert math.isclose(r["cpu_busy_pct"], 100 * (1 - 12 / 41), abs_tol=0.1)
    # iowait% = 100 * d_iowait(2) / d_total(41)
    assert math.isclose(r["cpu_iowait_pct"], 100 * (2 / 41), abs_tol=0.1)
    # mem used GB = (MemTotal-MemAvailable)/2**20
    assert math.isclose(r["mem_used_gb"], (2097152 - 943718) / 2**20, abs_tol=0.01)
    # cache GB = (Cached+SReclaimable)/2**20
    assert math.isclose(r["mem_cache_gb"], (210000 + 52000) / 2**20, abs_tol=0.01)
    # swap in MiB/s = d(pswpin=200)*4096/2**20 / dt(2)
    assert math.isclose(r["swap_in_mib_s"], 200 * 4096 / 2**20 / 2, abs_tol=0.01)
    # pressure page_scan_mib_s = 2048*4096/2**20 / 2 (reuse _compute_pressure_rates)
    assert math.isclose(r["page_scan_mib_s"], 2048 * 4096 / 2**20 / 2, abs_tol=0.01)
    # ublk aggregate cores: d(utime+stime)=200 /_CLK_TCK=100 / dt=2 -> 1.0
    assert math.isclose(r["ublk_cores"], 200 / 100 / 2, abs_tol=0.01)
    # first row -> NaN (no baseline)
    assert math.isnan(df.iloc[0]["cpu_busy_pct"])


def test_derive_host_guard_fires_on_vmstat_reset_and_cpu_gap(tmp_path, caplog):
    """A vmstat counter reset (cur<prev at a NORMAL interval) NaNs swap+pressure
    and emits a 'counter reset' WARNING; a host sampling gap (dt>MAX_INTERVAL_S)
    NaNs CPU and emits a 'gap detected' WARNING. Neither dilutes nor spikes.

    NOTE: the reset and gap are on SEPARATE rows because _guarded checks
    dt>max_interval (gap) BEFORE cur<prev (reset) -- a reset at a gap row is
    masked by the gap WARNING, so the reset path must be exercised at a
    normal-interval (dt<=MAX_INTERVAL_S) row to emit its own WARNING.
    """
    d = tmp_path
    # cpu: t0, t2, t4, t22. Rows at t=2,4 are clean (dt=2); t=22 is a GAP (dt=18>10).
    _write(
        d / "host_cpu_raw.csv",
        [
            "timestamp",
            "cpu_user",
            "cpu_nice",
            "cpu_system",
            "cpu_idle",
            "cpu_iowait",
            "cpu_irq",
            "cpu_softirq",
            "cpu_steal",
            "cpu_guest",
            "cpu_guest_nice",
            "ctxt",
            "btime",
            "processes",
            "procs_running",
            "procs_blocked",
            "softirq_total",
        ],
        [
            ["0", "100", "0", "50", "8000", "100", "10", "20", "5", "0", "0", "1", "0", "0", "0", "0", "0"],
            ["2", "120", "0", "55", "8010", "102", "11", "22", "6", "0", "0", "2", "0", "0", "0", "0", "0"],
            ["4", "130", "0", "60", "8020", "104", "12", "24", "7", "0", "0", "3", "0", "0", "0", "0", "0"],
            ["22", "140", "0", "65", "8030", "106", "13", "26", "8", "0", "0", "4", "0", "0", "0", "0", "0"],
        ],
    )
    _write(
        d / "host_mem_raw.csv",
        ["timestamp", "MemTotal", "MemAvailable", "MemFree", "Cached", "SReclaimable", "Buffers", "Dirty", "Writeback"],
        [
            ["0", "2097152", "1048576", "512000", "200000", "50000", "30000", "1000", "0"],
            ["2", "2097152", "943718", "450000", "210000", "52000", "32000", "2000", "0"],
            ["4", "2097152", "900000", "440000", "215000", "53000", "33000", "3000", "0"],
            ["22", "2097152", "880000", "430000", "218000", "54000", "34000", "4000", "0"],
        ],
    )
    # vmstat: pswpin 100 -> 300 -> 50 (RESET at t=4, dt=2 -> host_vmstat reset
    # guard fires, 'counter reset' WARNING) -> 60 at t=22 (dt=18 gap masks row4).
    _write(
        d / "host_vmstat_raw.csv",
        [
            "timestamp",
            "pswpin",
            "pswpout",
            "pgscan_kswapd",
            "pgscan_direct",
            "pgscan_direct_throttle",
            "pgsteal_kswapd",
            "pgsteal_direct",
            "workingset_refault_file",
        ],
        [
            ["0", "100", "0", "0", "0", "0", "0", "0", "0"],
            ["2", "300", "0", "2048", "0", "0", "0", "0", "0"],
            ["4", "50", "0", "3000", "0", "0", "0", "0", "0"],
            ["22", "60", "0", "3100", "0", "0", "0", "0", "0"],
        ],
    )

    with caplog.at_level(logging.WARNING):
        df = _derive_host(d)
    r2 = df.iloc[2]  # t=4: vmstat reset (dt=2, no gap) -> swap+pressure NaN
    r3 = df.iloc[3]  # t=22: cpu gap (dt=18>10) -> cpu_busy NaN
    # vmstat counter reset -> swap_in + page_scan NaN (NOT a negative spike, NOT 0)
    assert math.isnan(r2["swap_in_mib_s"])
    assert math.isnan(r2["page_scan_mib_s"])
    # cpu clean at t=4 (monotonic, dt=2) -> cpu_busy still computes
    assert not math.isnan(r2["cpu_busy_pct"])
    # cpu gap (dt=18>10) -> cpu_busy NaN
    assert math.isnan(r3["cpu_busy_pct"])
    assert any("gap detected" in m for m in caplog.messages)
    assert any("counter reset" in m and "host_vmstat" in m for m in caplog.messages)
    # row 1 (t=2, clean) still computes normally
    assert not math.isnan(df.iloc[1]["cpu_busy_pct"])
    assert not math.isnan(df.iloc[1]["swap_in_mib_s"])


def test_derive_per_vm_topk_and_other_and_reset(tmp_path, caplog):
    """per-VM CPU cores (delta) + PSS GB (gauge). VM count > MAX_VM_SERIES ->
    top-K by CPU busy mean + 'other' sum(skipna, min_count=1). A utime reset
    (pid restart, same vm_name) -> that point NaN, not a negative spike."""
    import logging

    d = tmp_path
    # 12 VMs: vm0 (highest busy, has the pid-restart reset) down to vm11 (lowest).
    # cores at t=2 = delta(utime+stime)/_CLK_TCK/dt; vm0=2000/100/2=10.0 (top),
    # vm1..vm7 step down, vm8..vm11 are the low-busy "other" set.
    rows = [["timestamp", "pid", "vm_name", "utime", "stime", "pss_mb"]]
    deltas = {0: 2000, 1: 1800, 2: 1600, 3: 1400, 4: 1200, 5: 1000, 6: 800, 7: 600, 8: 400, 9: 300, 10: 200, 11: 100}
    for v in range(12):
        rows.append(["0", f"{v}", f"vm{v}", "0", "0", f"{(v + 1) * 100}.0"])
        rows.append(["2", f"{v}", f"vm{v}", f"{deltas[v]}", "0", f"{(v + 1) * 100}.0"])
    # vm0 pid restart: same vm_name, utime drops 2000 -> 5 at t=4
    rows.append(["4", "999", "vm0", "5", "0", "50.0"])
    _write(d / "vm_cpu_raw.csv", rows[0], rows[1:])

    with caplog.at_level(logging.WARNING):
        cores, pss = _derive_per_vm(d, max_series=8)
    # top-8 = vm0..vm7 (named) + 'other' = vm8..vm11
    assert "other" in cores.columns
    named = [c for c in cores.columns if c.startswith("vm")]
    assert len(named) == 8
    assert "vm0" in named  # highest busy -> named, so its reset is observable
    # reset at t=4 (vm0 utime 2000->5): NaN, not a negative spike
    last = cores[cores["t_s"] == 4.0]
    assert math.isnan(last["vm0"].iloc[0])
    assert any("counter reset" in r and "vm0" in r for r in caplog.messages)
    # 'other' = sum(vm8..vm11 cores) at t=2 = (400+300+200+100)/_CLK_TCK/dt = 5.0
    other_t2 = cores[cores["t_s"] == 2.0]["other"].iloc[0]
    assert math.isclose(other_t2, (400 + 300 + 200 + 100) / 100 / 2, abs_tol=0.01)
    # PSS is a gauge (pss_mb/1024 -> GB); same column set as cores
    assert list(pss.columns) == list(cores.columns)
    pss_t2 = pss[pss["t_s"] == 2.0]["vm0"].iloc[0]
    assert math.isclose(pss_t2, 100.0 / 1024, abs_tol=0.001)  # vm0 pss_mb=100 at t=2


def test_derive_disk_rates_and_reset(tmp_path, caplog):
    """disk r/w MB/s + util% via base._compute_disk_io_rates; a sectors_read
    reset (device remount) -> that dev that interval NaN + WARNING."""
    import logging

    d = tmp_path
    from vm_monitor.base import _DISK_FIELDS

    hdr = ["timestamp", "device", *_DISK_FIELDS]

    def row(ts, dev, sr, sw, msio):
        vals = {f: 0 for f in _DISK_FIELDS}
        vals["sectors_read"] = sr
        vals["sectors_written"] = sw
        vals["ms_io"] = msio
        vals["reads_completed"] = sr
        vals["writes_completed"] = sw
        return [ts, dev, *[vals[f] for f in _DISK_FIELDS]]

    rows = [
        row("0", "sda", 0, 0, 0),
        row("2", "sda", 2000, 4000, 1000),
        row("4", "sda", 100, 100, 50),  # RESET: sectors_read 2000 -> 100
    ]
    _write(d / "disk_io_raw.csv", hdr, rows)

    with caplog.at_level(logging.WARNING):
        df = _derive_disk(d)
    # t=2: r_mb_s = 2000*512/2**20/2; _compute_disk_io_rates rounds to 2 dp -> 0.49
    r2 = df[df["t_s"] == 2.0].iloc[0]
    assert math.isclose(r2["sda_r_mb_s"], 2000 * 512 / 2**20 / 2, abs_tol=0.01)
    assert math.isclose(r2["sda_w_mb_s"], 4000 * 512 / 2**20 / 2, abs_tol=0.01)
    # t=4 reset -> NaN (not a negative spike)
    r4 = df[df["t_s"] == 4.0].iloc[0]
    assert math.isnan(r4["sda_r_mb_s"])
    assert any("counter reset" in m and "sda" in m for m in caplog.messages)


def test_derive_disk_inflight_gauge_excluded_from_guard(tmp_path, caplog):
    """inflight is a gauge (current I/Os in progress), not a cumulative counter;
    normal fluctuation (3->1->5) must NOT trip the reset guard. Cumulative
    counters stay monotonic -> every interval computes, no NaN holes on busy disks."""
    import logging

    d = tmp_path
    from vm_monitor.base import _DISK_FIELDS

    hdr = ["timestamp", "device", *_DISK_FIELDS]

    def row(ts, dev, sr, sw, msio, inflight):
        vals = {f: 0 for f in _DISK_FIELDS}
        vals["sectors_read"] = sr
        vals["sectors_written"] = sw
        vals["ms_io"] = msio
        vals["reads_completed"] = sr
        vals["writes_completed"] = sw
        vals["inflight"] = inflight  # gauge: fluctuates 3 -> 1 -> 5
        return [ts, dev, *[vals[f] for f in _DISK_FIELDS]]

    rows = [
        row("0", "sda", 0, 0, 0, 3),
        row("2", "sda", 2000, 4000, 1000, 1),  # inflight 3->1 (gauge drop); cumulative all up
        row("4", "sda", 4000, 8000, 2000, 5),  # inflight 1->5; cumulative all up
    ]
    _write(d / "disk_io_raw.csv", hdr, rows)

    with caplog.at_level(logging.WARNING):
        df = _derive_disk(d)
    # No false "counter reset" WARNING (inflight drop must not fire guard)
    assert not any("counter reset" in m and "sda" in m for m in caplog.messages)
    # Both intervals computed (not NaN)
    r2 = df[df["t_s"] == 2.0].iloc[0]
    r4 = df[df["t_s"] == 4.0].iloc[0]
    assert not math.isnan(r2["sda_r_mb_s"])
    assert not math.isnan(r4["sda_r_mb_s"])
    # And the rate is correct (proves it actually computed, not just non-NaN)
    assert math.isclose(r2["sda_r_mb_s"], 2000 * 512 / 2**20 / 2, abs_tol=0.01)


def _seed_all(d):
    """Minimal cumulative CSVs for all three sheets (host + per-VM + disk)."""
    _write(
        d / "host_cpu_raw.csv",
        [
            "timestamp",
            "cpu_user",
            "cpu_nice",
            "cpu_system",
            "cpu_idle",
            "cpu_iowait",
            "cpu_irq",
            "cpu_softirq",
            "cpu_steal",
            "cpu_guest",
            "cpu_guest_nice",
            "ctxt",
            "btime",
            "processes",
            "procs_running",
            "procs_blocked",
            "softirq_total",
        ],
        [
            ["0", "100", "0", "50", "8000", "100", "10", "20", "5", "0", "0", "1", "0", "0", "0", "0", "0"],
            ["2", "120", "0", "55", "8010", "102", "11", "22", "6", "0", "0", "2", "0", "0", "0", "0", "0"],
        ],
    )
    _write(
        d / "host_mem_raw.csv",
        ["timestamp", "MemTotal", "MemAvailable", "MemFree", "Cached", "SReclaimable", "Buffers", "Dirty", "Writeback"],
        [
            ["0", "2097152", "1048576", "512000", "200000", "50000", "30000", "1000", "0"],
            ["2", "2097152", "943718", "450000", "210000", "52000", "32000", "2000", "0"],
        ],
    )
    _write(
        d / "host_vmstat_raw.csv",
        [
            "timestamp",
            "pswpin",
            "pswpout",
            "pgscan_kswapd",
            "pgscan_direct",
            "pgscan_direct_throttle",
            "pgsteal_kswapd",
            "pgsteal_direct",
            "workingset_refault_file",
        ],
        [["0", "100", "0", "0", "0", "0", "0", "0", "0"], ["2", "300", "0", "2048", "0", "0", "0", "0", "0"]],
    )
    _write(
        d / "ublk_cpu_raw.csv",
        ["timestamp", "pid", "utime", "stime"],
        [["0", "4242", "0", "0"], ["2", "4242", "100", "100"]],
    )
    _write(
        d / "vm_cpu_raw.csv",
        ["timestamp", "pid", "vm_name", "utime", "stime", "pss_mb"],
        [["0", "1", "vm0", "0", "0", "1024.0"], ["2", "1", "vm0", "200", "0", "1024.0"]],
    )
    from vm_monitor.base import _DISK_FIELDS

    zero = [0] * len(_DISK_FIELDS)
    row2 = [0] * len(_DISK_FIELDS)
    row2[_DISK_FIELDS.index("sectors_read")] = 2000
    _write(d / "disk_io_raw.csv", ["timestamp", "device", *_DISK_FIELDS], [["0", "sda", *zero], ["2", "sda", *row2]])


def test_build_xlsx_sheets_and_charts(tmp_path):
    d = tmp_path
    _seed_all(d)
    out = d / "host_resources.xlsx"
    build_host_resources_xlsx(d, out)
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {"Host resources", "Per-VM", "Disk IO"}
    host = wb["Host resources"]
    assert len(host._charts) == 5
    # series counts lock the col-set per chart (memory=4, cpu=2, pressure=3,
    # swap=2, cores=2) -- guards against pressure wrongly including swap cols
    # or an off-by-one skipping a series.
    assert sorted(len(c.series) for c in host._charts) == [2, 2, 2, 3, 4]
    assert len(wb["Per-VM"]._charts) == 2
    assert len(wb["Disk IO"]._charts) == 2


def test_build_xlsx_degrades_on_missing_csv(tmp_path):
    # only host_cpu_raw present -> only Host resources sheet, no crash
    _write(
        tmp_path / "host_cpu_raw.csv",
        [
            "timestamp",
            "cpu_user",
            "cpu_nice",
            "cpu_system",
            "cpu_idle",
            "cpu_iowait",
            "cpu_irq",
            "cpu_softirq",
            "cpu_steal",
            "cpu_guest",
            "cpu_guest_nice",
            "ctxt",
            "btime",
            "processes",
            "procs_running",
            "procs_blocked",
            "softirq_total",
        ],
        [["0", "1", "0", "1", "90", "0", "0", "0", "0", "0", "0", "1", "0", "0", "0", "0", "0"]],
    )
    out = tmp_path / "host_resources.xlsx"
    build_host_resources_xlsx(tmp_path, out)
    wb = openpyxl.load_workbook(out)
    assert "Host resources" in wb.sheetnames
    assert "Per-VM" not in wb.sheetnames
    assert "Disk IO" not in wb.sheetnames


def test_build_xlsx_no_crash_when_all_csvs_missing(tmp_path, caplog):
    """All raw CSVs missing -> no sheet built -> WARNING + no file written,
    never raises (contract: 'never raises'). Regression for the wb.save()
    IndexError when wb.remove(wb.active) left a 0-sheet workbook."""
    out = tmp_path / "host_resources.xlsx"
    with caplog.at_level(logging.WARNING):
        result = build_host_resources_xlsx(tmp_path, out)
    assert result == out  # returns the intended path without raising
    assert not out.exists()  # no misleading empty workbook written
    assert any("no raw CSV data" in m for m in caplog.messages)
